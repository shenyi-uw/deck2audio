#!/usr/bin/env python3
"""deck2audio — turn a PowerPoint deck into a narrated MP3, with a human-in-the-loop LLM step.

Stages (each resumable; state lives in the project folder):
  1. extract   Pull slide text + speaker notes from a .pptx
  2. research  Generate LLM prompts for academic sources & real-world examples; you paste back + approve
  3. script    Draft with an LLM or use deck content directly; review + approve
  4. tts       Synthesize the approved script with a neural female voice and export an MP3

Run everything:  python deck2audio.py run deck.pptx
"""
from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path
from typing import Callable

import kokoro_tts
import macos_tts
import pcm_audio

DEFAULT_VOICE = "en-US-AvaMultilingualNeural"
WORDS_PER_MINUTE = 150
SLIDE_HEADING = re.compile(r"^\s*#{1,6}\s*Slide\s+(\d+)\b.*$", re.IGNORECASE | re.MULTILINE)
VOICE_PREVIEW_TEXT = (
    "Hi there! I'm going to be your narrator today. "
    "Grab a coffee, get comfortable, and let's dive into some genuinely fascinating ideas together."
)


# ----------------------------------------------------------------------------- helpers

def project_dir_for(deck: Path) -> Path:
    return deck.with_suffix("").parent / f"{deck.stem}_narration"


def load_state(proj: Path) -> dict:
    f = proj / "state.json"
    return json.loads(f.read_text()) if f.exists() else {}


def save_state(proj: Path, state: dict) -> None:
    (proj / "state.json").write_text(json.dumps(state, indent=2))


def load_slides(proj: Path) -> list[dict]:
    f = proj / "slides.json"
    if not f.exists():
        sys.exit(f"No slides.json in {proj}. Run the 'extract' stage first.")
    return json.loads(f.read_text())


def banner(title: str) -> None:
    print("\n" + "=" * 78 + f"\n  {title}\n" + "=" * 78)


def ask(question: str, choices: str) -> str:
    """Single-key style prompt; `choices` like 'ARES' — first letter is default."""
    opts = "/".join(c if i else c.upper() for i, c in enumerate(choices.lower()))
    while True:
        ans = input(f"{question} [{opts}]: ").strip().lower()
        if not ans:
            return choices[0].lower()
        if ans[0] in choices.lower():
            return ans[0]


def copy_to_clipboard(text: str) -> bool:
    for cmd in (["pbcopy"], ["wl-copy"], ["xclip", "-selection", "clipboard"], ["clip"]):
        if shutil.which(cmd[0]):
            subprocess.run(cmd, input=text.encode(), check=False)
            return True
    return False


def read_clipboard() -> str | None:
    for cmd in (["pbpaste"], ["wl-paste"], ["xclip", "-selection", "clipboard", "-o"]):
        if shutil.which(cmd[0]):
            return subprocess.run(cmd, capture_output=True, check=False).stdout.decode()
    return None


def edit_in_editor(text: str, suffix: str = ".md") -> str:
    editor = os.environ.get("VISUAL") or os.environ.get("EDITOR") or ("nano" if shutil.which("nano") else "vi")
    with tempfile.NamedTemporaryFile("w", suffix=suffix, delete=False) as tmp:
        tmp.write(text)
        path = tmp.name
    subprocess.call([*editor.split(), path])
    out = Path(path).read_text()
    os.unlink(path)
    return out


def collect_response() -> str:
    """Get an LLM response from the user.

    Clipboard is the default because terminals (esp. macOS) truncate very long pasted lines.
    """
    print(textwrap.dedent("""
        How do you want to provide the LLM's response?
          [Enter]  read it from the clipboard (copy the full response first)
          e        open your $EDITOR and paste it there
          p        paste directly here, then type END on its own line
          f        load it from a file path
    """))
    while True:
        mode = input("> ").strip().lower()
        if mode in ("", "c"):
            text = read_clipboard()
            if text is None:
                print("No clipboard tool available; choose e, p, or f.")
                continue
        elif mode == "e":
            text = edit_in_editor("<!-- Paste the LLM response below, save, and close. -->\n\n")
            text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
        elif mode == "p":
            lines = []
            try:
                while (line := input()) .strip() != "END":
                    lines.append(line)
            except EOFError:
                pass
            text = "\n".join(lines)
        elif mode == "f":
            p = Path(input("File path: ").strip().strip("'\"")).expanduser()
            if not p.exists():
                print("File not found.")
                continue
            text = p.read_text()
        else:
            continue
        if text.strip():
            return text.strip()
        print("Got an empty response — try again.")


def human_in_the_loop(prompt: str, label: str, validate=None) -> str | None:
    """Show a prompt, copy it, wait for the pasted response, and require explicit approval.

    Returns the approved text, or None if the user skips this step.
    """
    banner(label)
    print(prompt)
    print("-" * 78)
    if copy_to_clipboard(prompt):
        print("✓ The prompt above has been copied to your clipboard.")
    print("Paste it into your LLM of choice (ChatGPT, Claude, Copilot, Gemini, …).")
    if ask("Ready to provide the response, or skip this step?", "rs") == "s":
        return None

    response = collect_response()
    while True:
        problems = validate(response) if validate else []
        banner(f"Review — {label}")
        print(response)
        print("-" * 78)
        print(f"{len(response.split())} words.")
        for p in problems:
            print(f"⚠  {p}")
        choice = ask("[A]pprove, [E]dit, [R]e-paste a new response, [S]kip?", "aers")
        if choice == "a":
            if problems and ask("There are warnings above. Approve anyway?", "ny") == "n":
                continue
            return response
        if choice == "e":
            response = edit_in_editor(response).strip()
        elif choice == "r":
            print("\n(Tip: ask the LLM to revise, then copy its new answer.)")
            response = collect_response()
        else:
            return None


# ----------------------------------------------------------------------------- stage 1: extract

def shape_text(shape) -> list[str]:
    out: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip()
            if t:
                out.append(("  " * para.level) + t)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            out.extend(shape_text(s))
    return out


def extract(deck: Path, proj: Path) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(str(deck))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        body: list[str] = []
        for shape in slide.shapes:
            if slide.shapes.title is not None and shape.shape_id == slide.shapes.title.shape_id:
                continue
            body.extend(shape_text(shape))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"number": i, "title": title, "body": body, "notes": notes})

    proj.mkdir(parents=True, exist_ok=True)
    (proj / "slides.json").write_text(json.dumps(slides, indent=2, ensure_ascii=False))
    (proj / "slides.md").write_text(slides_as_markdown(slides))
    state = load_state(proj)
    state.update({"deck": str(deck.resolve()), "extracted": True})
    save_state(proj, state)
    print(f"✓ Extracted {len(slides)} slides → {proj / 'slides.md'}")
    return slides


def slides_as_markdown(slides: list[dict]) -> str:
    parts = []
    for s in slides:
        parts.append(f"### Slide {s['number']}: {s['title'] or '(untitled)'}")
        if s["body"]:
            parts.append("Slide text:\n" + "\n".join(f"- {b}" for b in s["body"]))
        parts.append("Speaker notes:\n" + (s["notes"] or "(none)"))
        parts.append("")
    return "\n".join(parts)


# ----------------------------------------------------------------------------- stage 2: research

RESEARCH_PROMPTS = {
    "academic": """\
You are a research assistant helping me prepare a narrated presentation.
Below is the full content of my slide deck (slide text + speaker notes).
Intended audience: {audience}. Adapt the explanation depth to this audience.

TASK: Suggest academic grounding that would make the narration more insightful.
For the slides where it genuinely adds value (not every slide), provide:
  • Key theories, frameworks, or seminal studies (author(s), year, venue)
  • One or two sentences on the core finding and HOW it connects to the slide
  • For each specific study, the research team (including the lead researcher),
    the institution at the time of the study, and the year of publication.
    Verify these details; mark unknown details as "verify", never guess them.
    Describe what the study actually found, its scope, and relevant limitations.
  • A short, quotable takeaway phrased for a spoken audience

RULES:
  • Only cite sources you are confident exist. If unsure, say "verify" next to it.
  • No fabricated statistics. Prefer well-known, verifiable work.
  • Organize your answer by "Slide N" headings so it maps back to the deck.
  • Keep it concise: at most ~3 items per slide.
  • Recommend spoken source attribution only when describing a specific study's
    finding, not to decorate a general statement with a researcher's name.
    Do not assume the audience knows a researcher; introduce the team and institution.

DECK:
{deck}
""",
    "examples": """\
You are helping me make a presentation memorable for {audience}.
Below is the full content of my slide deck (slide text + speaker notes).

TASK: For the slides where it fits, suggest vivid real-world examples, case studies,
analogies, or short stories that illustrate the point. For each:
  • A 2–4 sentence description told in a way that works when spoken aloud
  • Why it illustrates the slide's idea
  • Approximate date/company/person, so I can verify it

RULES:
  • Use real, verifiable events; flag anything uncertain with "verify".
  • Favor diverse industries and recent examples (last ~10 years) where possible.
  • Choose examples relatable to this audience's experience, age, and knowledge.
    Explain unfamiliar contexts without condescension. Label invented analogies
    as hypothetical, never as actual events.
  • Organize your answer by "Slide N" headings.
  • At most ~2 examples per slide; skip slides where an example would feel forced.
  • Plan examples across the ENTIRE presentation. Assign each example, case study,
    story, or analogy to its best-fitting slide rather than suggesting it repeatedly.
    If several slides share an idea, use distinct examples or explain it without
    another example. Do not repeat an example already supplied in the deck.

DECK:
{deck}
""",
}


def research_prompt(key: str, slides: list[dict], audience: str = "") -> str:
    return RESEARCH_PROMPTS[key].format(
        deck=slides_as_markdown(slides), audience=audience or "a general professional audience")


def research(proj: Path, audience: str = "") -> None:
    slides = load_slides(proj)
    state = load_state(proj)
    rdir = proj / "research"
    rdir.mkdir(exist_ok=True)
    labels = {"academic": "Academic resources & theory", "examples": "Real-world examples & stories"}

    for key in RESEARCH_PROMPTS:
        out = rdir / f"{key}.md"
        status = state.get("research", {}).get(key)
        if status in ("approved", "skipped"):
            if ask(f"'{labels[key]}' is already {status}. Redo it?", "ny") == "n":
                continue
        prompt = research_prompt(key, slides, audience)
        (rdir / f"{key}.prompt.txt").write_text(prompt)
        result = human_in_the_loop(prompt, f"Step 2 · {labels[key]}")
        state.setdefault("research", {})[key] = "approved" if result else "skipped"
        if result:
            out.write_text(result)
            print(f"✓ Saved approved research → {out}")
        elif out.exists():
            out.unlink()
        save_state(proj, state)


def approved_research(proj: Path) -> str:
    state = load_state(proj)
    parts = []
    for key, title in (("academic", "ACADEMIC RESOURCES"), ("examples", "REAL-WORLD EXAMPLES")):
        f = proj / "research" / f"{key}.md"
        if state.get("research", {}).get(key) == "approved" and f.exists():
            parts.append(f"=== APPROVED {title} (use where they genuinely help) ===\n{f.read_text()}")
    return "\n\n".join(parts)


# ----------------------------------------------------------------------------- stage 3: script

SCRIPT_PROMPT = """\
You are an award-winning keynote speaker and speechwriter. Write the word-for-word
NARRATION SCRIPT for a recorded audio presentation of the slide deck below.
It will be read aloud by a text-to-speech narrator{audience}.

TONE & STYLE
  • Wise, charismatic, insightful, and personable — like a favorite teacher who is
    also a great storyteller. Be engaging, not pompous or preachy.
  • Use occasional gentle, appropriate humor where it fits; never force jokes,
    mock people, trivialize sensitive topics, or distract from the message.
  • Adapt vocabulary, explanation depth, and real-world examples to the intended
    audience's age, knowledge, and everyday experience without condescension.
  • Sound like a personable one-person podcast, not a formal lecture or a panel.
    Address one listener as "you" and use "I" naturally as the narrator
    (e.g., "I want to show you..." or "Here's how I'd think about it").
    Use "we" when it genuinely invites shared thinking, not as a substitute for
    direct conversation. Use contractions, natural rhythm, brief reflective
    questions, and warm transitions without mechanically repeating verbal hooks.
    Do not invent personal experiences, credentials, or anecdotes for the narrator.
  • Open with a hook (a question, surprising fact, or short story){intro_note}.
  • Don't just read the bullets — explain WHY they matter, connect ideas, and add insight.
  • Weave in the approved academic sources and real-world examples where they genuinely
    strengthen a point; follow the spoken-source rules below, not name-dropping.
  • Add concise, relatable real-world examples where helpful even if no separate
    examples research was supplied. Use verifiable events or clearly hypothetical
    analogies; do not fabricate case studies, names, dates, or statistics.
  • Avoid repeated examples across the ENTIRE presentation, including earlier batches.
    Use each case study, story, or analogy once, on the slide where it helps most.
    Check the accepted earlier narration below before introducing an example.
    A brief, clearly signposted callback is fine only if it adds a new connection;
    do not retell the example or present it as new. In the closing summary, recap
    takeaways rather than telling the same examples again.
  • Smooth spoken transitions between slides. {outro_note}
  • Target roughly {wpm_low}–{wpm_high} words per slide on average, more for dense slides,
    less for title/section slides.{duration_note}

SPOKEN SOURCES & STUDIES
  • Make reference to a source only when describing a specific study's finding,
    explaining what was studied and found and why that finding matters here.
    Do not attach a researcher's name to an ordinary statement, broad idea, or
    opinion. Do not assume the listener recognizes researchers' names.
  • When describing a specific study, introduce the research team (including the
    lead researcher), the institution at the time of the study, and the year of
    publication in a concise, conversational sentence alongside the finding.
    Verify the attribution, institution, year, and finding; never invent these
    details or infer an institution from a researcher's current affiliation.
  • If these study details cannot be verified, omit that specific study reference.
    Only retain a general explanation if it is independently supportable and
    carefully worded; otherwise omit the unsupported claim as well.
  • The exception for mentioning a person outside a study is a commonly known
    public or historical figure when genuinely relevant to the message. This
    is not permission to name-drop, fabricate quotations, or substitute fame
    for evidence. Avoid academic roll calls and unexplained "according to" lines.

FACT-CHECKING & FAITHFULNESS
  • Fact-check claims in the existing slide text AND speaker notes before drafting.
    Check dates, statistics, attributions, causal claims, and sweeping generalizations
    against reliable sources when you have source access. Approved research is not
    automatically proof; check its claims too.
  • Do not pretend to have verified anything you cannot verify. If source access
    is unavailable or evidence is uncertain, omit unsupported specifics or use
    appropriately cautious wording; never invent sources or certainty.
  • Avoid over-claiming: distinguish evidence from opinion, correlation from
    causation, and possibilities from established results. Preserve important caveats.
  • The goal is to promote the deck's original message with appropriate clarifications
    and careful wording, NOT to critique or review the deck. Do not repeat known false
    claims; correct or reframe them naturally without attacking the author or adding
    a critique section. Do not endorse a false premise merely to preserve the message.
  • Output only the spoken script, not a fact-check report. Human verification is
    still required for consequential claims; do not describe this draft as certified.

FORMAT — THIS IS PARSED BY A PROGRAM, FOLLOW EXACTLY
  • One section per slide, in order, each starting with a heading line exactly like:
        ## Slide 7
  • Include EVERY slide from {first} to {last}, even if a section is one sentence.
  • Under each heading write ONLY the words to be spoken: plain prose, no bullet points,
    no markdown formatting, no URLs, no stage directions — except you MAY insert
    [pause] on its own where a dramatic beat helps.
  • Spell out symbols and abbreviations the way they should be spoken
    (e.g., "50 percent", "versus", "A I" only if it should be spelled out).
  • Do not add any text before the first heading or after the last section.
{continuity}
{research}
{summary_context}

=== SLIDE DECK{batch_note} ===
{deck}
"""


def parse_script(text: str) -> dict[int, str]:
    sections: dict[int, str] = {}
    matches = list(SLIDE_HEADING.finditer(text))
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        sections[int(m.group(1))] = text[m.end():end].strip()
    return sections


def validate_script(text: str, expected: set[int]) -> list[str]:
    got = parse_script(text)
    problems: list[str] = []
    if missing := sorted(expected - got.keys()):
        problems.append(f"Missing sections for slides: {missing}")
    if extra := sorted(got.keys() - expected):
        problems.append(f"Unexpected slide numbers: {extra}")
    if empty := sorted(k for k, v in got.items() if k in expected
                       and not re.sub(r"\[pause\]", "", clean_for_speech(v), flags=re.I).strip()):
        problems.append(f"Empty sections: {empty}")
    numbers = [int(m.group(1)) for m in SLIDE_HEADING.finditer(text)]
    if len(numbers) != len(set(numbers)):
        problems.append("Duplicate slide headings.")
    if numbers != sorted(numbers):
        problems.append("Slide sections must be in deck order.")
    if matches := list(SLIDE_HEADING.finditer(text)):
        if text[:matches[0].start()].strip():
            problems.append("Remove text before the first slide heading.")
    return problems


def format_script(slides: list[dict], sections: dict[int, str]) -> str:
    return "\n\n".join(
        f"## Slide {s['number']}: {s['title'] or ''}".rstrip(": ") + f"\n\n{sections.get(s['number'], '')}"
        for s in slides)


def script_prompt(slides: list[dict], batch_size: int, batch_index: int, audience: str,
                  minutes: float | None, sections: dict[int, str], research_text: str) -> str:
    if not slides or batch_size < 1:
        raise ValueError("Script drafting requires slides and a positive batch size.")
    if minutes is not None and minutes <= 0:
        raise ValueError("Target duration must be positive.")
    batches = [slides[i:i + batch_size] for i in range(0, len(slides), batch_size)]
    if not 0 <= batch_index < len(batches):
        raise ValueError("Script batch index is out of range.")
    batch = batches[batch_index]
    first, last = batch[0]["number"], batch[-1]["number"]
    is_first, is_last = batch_index == 0, batch_index == len(batches) - 1
    total_words = int(minutes * WORDS_PER_MINUTE) if minutes else None
    per_slide = (total_words // len(slides)) if total_words else 110
    continuity = ""
    if not is_first:
        prev_num = slides[batch_index * batch_size - 1]["number"]
        prior_slides = [s for s in slides[:batch_index * batch_size] if s["number"] in sections]
        continuity = (
            f"\nCONTINUITY: This is part {batch_index + 1} of {len(batches)}. The narration so far ended "
            f"with slide {prev_num}:\n\"\"\"\n{sections.get(prev_num, '')[-1200:]}\n\"\"\"\n"
            "Continue seamlessly — do NOT re-introduce the talk or greet the listener again.\n"
            "\n=== ACCEPTED EARLIER NARRATION (reference only; do not rewrite or repeat) ===\n"
            "Check ALL earlier sections for examples already used and for continuity. "
            "Use different examples in this batch, or no example where none is needed.\n"
            f"{format_script(prior_slides, sections)}\n"
        )
    return SCRIPT_PROMPT.format(
        audience=f" for {audience or 'a general professional audience'}",
        intro_note=" and briefly preview the journey" if is_first else " — only in part 1; this is a continuation",
        outro_note=(
            "At the END of the LAST slide's narration, summarize the ENTIRE presentation's "
            "key ideas and takeaways, whether or not the deck contains a summary slide. "
            "Include earlier slides, not just this batch. Finish with a memorable close and "
            "an appropriate call to action. Keep the summary inside the last slide's section; "
            "do not add a new slide or text outside the slide sections."
            if is_last else "Do NOT conclude or summarize the whole talk yet; more slides follow this batch."),
        wpm_low=max(40, per_slide - 30), wpm_high=per_slide + 40,
        duration_note=(f" The full talk should run about {minutes:g} minutes (~{total_words} words)."
                       if minutes else ""),
        first=first, last=last, continuity=continuity,
        research=("\n" + research_text) if research_text else "",
        summary_context=(
            "\n=== FULL DECK CONTEXT FOR THE CLOSING SUMMARY (reference only; narrate ONLY this batch) ===\n"
            + slides_as_markdown(slides) if is_last and len(batches) > 1 else ""),
        batch_note=f" (slides {first}–{last} of {len(slides)})" if len(batches) > 1 else "",
        deck=slides_as_markdown(batch),
    )


def choose_skip_all() -> bool:
    return ask("Draft with an [L]LM, or [S]kip all research and LLM prompts and use deck content?", "ls") == "s"


def deck_script_sections(slides: list[dict]) -> dict[int, str]:
    sections: dict[int, str] = {}
    for slide in slides:
        notes = clean_for_speech(slide["notes"])
        text = notes or clean_for_speech("\n\n".join([slide["title"], *slide["body"]]))
        if not text:
            sys.exit(f"Slide {slide['number']} has no usable text or speaker notes. Add speaker notes before drafting.")
        sections[slide["number"]] = text
    return sections


def script(proj: Path, batch_size: int, audience: str, minutes: float | None,
           skip_all: bool | None = None) -> None:
    slides = load_slides(proj)
    state = load_state(proj)
    script_file = proj / "script.md"
    if state.get("script") == "approved" and script_file.exists():
        if ask("An approved script already exists. Redo it?", "ny") == "n":
            return

    if not slides:
        sys.exit("No slides found. Extract a deck with slides before drafting.")
    if skip_all is None:
        skip_all = choose_skip_all()
    if skip_all:
        print("Using speaker notes, falling back to slide titles and text. No research or LLM prompts.")
        if audience or minutes:
            print("Audience and target duration apply only to LLM drafting; deck content is used as-is.")
        approve_script(proj, slides, deck_script_sections(slides))
        return

    research_text = approved_research(proj)
    if batch_size < 1:
        sys.exit("Slides per script prompt must be at least one.")
    n = len(slides)
    batches = [slides[i:i + batch_size] for i in range(0, n, batch_size)]
    sections: dict[int, str] = {}

    for bi, batch in enumerate(batches):
        first, last = batch[0]["number"], batch[-1]["number"]
        prompt = script_prompt(slides, batch_size, bi, audience, minutes, sections, research_text)
        (proj / f"script_prompt_{bi + 1}.txt").write_text(prompt)
        expected = {s["number"] for s in batch}

        def validate(text: str, expected=expected) -> list[str]:
            return validate_script(text, expected)

        label = f"Step 3 · Narration script" + (f" — part {bi + 1}/{len(batches)} (slides {first}–{last})"
                                                 if len(batches) > 1 else "")
        result = human_in_the_loop(prompt, label, validate)
        if result is None:
            print("Script drafting cancelled; nothing saved.")
            return
        sections.update({k: v for k, v in parse_script(result).items() if k in expected})

    approve_script(proj, slides, sections)


def approve_script(proj: Path, slides: list[dict], sections: dict[int, str]) -> None:
    full = format_script(slides, sections)
    words = len(full.split())
    banner("Final review — full narration script")
    print(full)
    print("-" * 78)
    print(f"{words} words ≈ {words / WORDS_PER_MINUTE:.1f} minutes of audio.")
    while True:
        c = ask("Confirm this script for audio production? [A]pprove, [E]dit, [C]ancel", "aec")
        if c == "a":
            break
        if c == "e":
            full = edit_in_editor(full)
            print(f"{len(full.split())} words ≈ {len(full.split()) / WORDS_PER_MINUTE:.1f} minutes.")
        else:
            print("Not approved; draft saved to script.draft.md")
            (proj / "script.draft.md").write_text(full)
            state = load_state(proj)
            state["script"] = "draft"
            save_state(proj, state)
            return
    script_file = proj / "script.md"
    script_file.write_text(full)
    state = load_state(proj)
    state["script"] = "approved"
    save_state(proj, state)
    print(f"✓ Approved script saved → {script_file}  (you can still hand-edit it before TTS)")


# ----------------------------------------------------------------------------- stage 4: TTS

def clean_for_speech(text: str) -> str:
    text = re.sub(r"<!--.*?-->", " ", text, flags=re.S)
    text = re.sub(r"https?://\S+", "", text)
    text = re.sub(r"\[(?!pause\])([^\]]+)\]\([^)]+\)", r"\1", text, flags=re.I)  # md links
    text = re.sub(r"\[(\d+(,\s*\d+)*)\]", "", text)                                 # [1] citations
    text = re.sub(r"^\s*#+\s*", "", text, flags=re.M)
    text = re.sub(r"^\s*([-*•]|\d+\.)\s+", "", text, flags=re.M)
    text = re.sub(r"[*_`~]+", "", text)
    text = text.replace("&", " and ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s+([.,;:!?])", r"\1", text)
    return text.strip()


def silent_mp3(seconds: float, header: bytes) -> bytes:
    """Build silent MPEG audio frames matching the encoder's own frame header."""
    versions = {0b11: 1, 0b10: 2, 0b00: 2.5}
    b1, b2, b3 = header[1], header[2], header[3]
    ver = versions[(b1 >> 3) & 0b11]
    br_idx, sr_idx = b2 >> 4, (b2 >> 2) & 0b11
    br_table = ([0, 32, 40, 48, 56, 64, 80, 96, 112, 128, 160, 192, 224, 256, 320] if ver == 1
                else [0, 8, 16, 24, 32, 40, 48, 56, 64, 80, 96, 112, 128, 144, 160])
    sr_table = {1: [44100, 48000, 32000], 2: [22050, 24000, 16000], 2.5: [11025, 12000, 8000]}[ver]
    bitrate, srate = br_table[br_idx] * 1000, sr_table[sr_idx]
    samples = 1152 if ver == 1 else 576
    size = (144 if ver == 1 else 72) * bitrate // srate
    frame = bytes([0xFF, b1 | 0x01, b2 & 0xFD, b3]) + bytes(size - 4)  # no CRC, no padding
    return frame * max(1, round(seconds * srate / samples))


async def synth_chunk(text: str, voice: str, rate: str, pitch: str, retries: int = 3,
                      *, backend: str = "edge") -> bytes:
    if backend in ("macos", "kokoro"):
        engine = kokoro_tts if backend == "kokoro" else macos_tts
        return pcm_audio.encode_mp3(await engine.synth_pcm(text, voice, rate, pitch))
    if backend != "edge":
        raise ValueError(f"Unknown speech backend: {backend}")
    import edge_tts

    for attempt in range(retries):
        try:
            buf = bytearray()
            async for chunk in edge_tts.Communicate(text, voice, rate=rate, pitch=pitch).stream():
                if chunk["type"] == "audio":
                    if "data" not in chunk:
                        raise RuntimeError("TTS audio chunk is missing data")
                    buf.extend(chunk["data"])
            if buf:
                return bytes(buf)
        except Exception as e:  # network hiccups are common; retry
            if attempt == retries - 1:
                raise
            print(f"   retrying after error: {e}")
        await asyncio.sleep(2 * (attempt + 1))
    raise RuntimeError("TTS returned no audio")


async def tts_async(proj: Path, voice: str, rate: str, pitch: str, slide_gap: float, pause_len: float,
                    output: Path, *, backend: str = "edge") -> None:
    script_file = proj / "script.md"
    if not script_file.exists():
        sys.exit("No script.md found. Run the 'script' stage first.")
    state = load_state(proj)
    if state.get("script") != "approved":
        if ask("script.md has not been approved through the workflow. Use it anyway?", "ny") == "n":
            return
    sections = parse_script(script_file.read_text())
    if not sections:
        sys.exit("Couldn't find any '## Slide N' sections in script.md.")

    words = sum(len(v.split()) for v in sections.values())
    banner("Step 4 · Text-to-speech")
    print(f"Backend: {backend}   voice: {voice or 'System default'}   rate: {rate}   pitch: {pitch}")
    print(f"{len(sections)} slides, {words} words ≈ {words / WORDS_PER_MINUTE:.1f} min")
    if ask("Generate the audio now?", "yn") == "n":
        return

    await generate_audio(proj, sections, voice, rate, pitch, slide_gap, pause_len, output, backend=backend)


async def generate_audio(proj: Path, sections: dict[int, str], voice: str, rate: str, pitch: str,
                         slide_gap: float, pause_len: float, output: Path,
                         progress: Callable[[int, int], None] | None = None,
                         *, backend: str = "edge") -> None:
    if backend not in ("edge", "macos", "kokoro"):
        raise ValueError(f"Unknown speech backend: {backend}")
    if not sections or any(not re.sub(r"\[pause\]", "", clean_for_speech(text), flags=re.I).strip()
                           for text in sections.values()):
        raise ValueError("Every slide must have spoken text before generating audio.")
    if slide_gap < 0 or pause_len < 0:
        raise ValueError("Slide gaps and pauses cannot be negative.")
    local = backend != "edge"
    engine = kokoro_tts if backend == "kokoro" else macos_tts
    if local:
        engine.validate_settings(voice, rate, pitch)
    adir = proj / "audio"
    adir.mkdir(exist_ok=True)
    pieces: list[bytes] = []
    header = None
    for index, num in enumerate(sorted(sections), 1):
        text = clean_for_speech(sections[num])
        segments = [s.strip() for s in re.split(r"\[pause\]", text, flags=re.I)]
        slide_audio: list[bytes] = []
        print(f" • Slide {num} ({len(text.split())} words)")
        for si, seg in enumerate(segments):
            if seg:
                audio = (await engine.synth_pcm(seg, voice, rate, pitch) if local
                         else await synth_chunk(seg, voice, rate, pitch))
                header = header or audio[:4]
                slide_audio.append(audio)
            if si < len(segments) - 1 and header and pause_len > 0:
                slide_audio.append(pcm_audio.silence(pause_len) if local else silent_mp3(pause_len, header))
        blob = b"".join(slide_audio)
        (adir / f"slide_{num:03d}.mp3").write_bytes(pcm_audio.encode_mp3(blob) if local else blob)
        if pieces and header and slide_gap > 0:
            pieces.append(pcm_audio.silence(slide_gap) if local else silent_mp3(slide_gap, header))
        pieces.append(blob)
        if progress:
            progress(index, len(sections))

    output.parent.mkdir(parents=True, exist_ok=True)
    combined = b"".join(pieces)
    output.write_bytes(pcm_audio.encode_mp3(combined) if local else combined)
    state = load_state(proj)
    state["audio"] = str(output)
    save_state(proj, state)
    print(f"\n✓ Narrated presentation → {output}")
    print(f"  Per-slide clips → {adir}/ (handy for embedding audio into each slide)")


def tts(proj: Path, args) -> None:
    out = Path(args.output) if args.output else proj / f"{Path(load_state(proj).get('deck', 'presentation')).stem}_narrated.mp3"
    voice = (args.macos_voice if args.backend == "macos" else
             args.kokoro_voice if args.backend == "kokoro" else args.voice)
    asyncio.run(tts_async(proj, voice, args.rate, args.pitch, args.slide_gap, args.pause, out,
                          backend=args.backend))


async def preview_voice(voice: str, rate: str, pitch: str, out: Path, *, backend: str = "edge") -> None:
    out.write_bytes(await synth_chunk(VOICE_PREVIEW_TEXT, voice, rate, pitch, backend=backend))
    print(f"✓ Voice sample → {out}")
    if shutil.which("afplay"):
        subprocess.call(["afplay", str(out)])


# ----------------------------------------------------------------------------- CLI

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("ui", help="open the local browser interface")

    def add_tts_opts(p):
        p.add_argument("--backend", choices=("edge", "macos", "kokoro"), default="edge",
                       help="speech engine: edge (online), macos or kokoro (offline)")
        p.add_argument("--kokoro-voice", choices=tuple(kokoro_tts.VOICES), default=kokoro_tts.DEFAULT_VOICE,
                       help="American-English Kokoro voice (default: af_heart)")
        p.add_argument("--macos-voice", default="",
                       help="installed macOS voice name; default is the system voice (list: say -v '?')")
        p.add_argument("--voice", default=DEFAULT_VOICE,
                       help=f"edge-tts voice (default {DEFAULT_VOICE}; try en-US-EmmaMultilingualNeural, "
                            "en-US-JennyNeural, en-GB-SoniaNeural)")
        p.add_argument("--rate", default="+0%", help="speaking rate, e.g. -5%% or +8%%")
        p.add_argument("--pitch", default="+0Hz", help="pitch shift, e.g. +2Hz")

    def add_proj(p):
        p.add_argument("target", help="the .pptx file or its *_narration project folder")

    p = sub.add_parser("run", help="run the full workflow (resumes where you left off)")
    add_proj(p)
    p.add_argument("--audience", default="", help='e.g. "graduate students in public health"')
    p.add_argument("--minutes", type=float, help="target talk length in minutes")
    p.add_argument("--batch-size", type=int, default=15, help="slides per script prompt (large decks)")
    p.add_argument("--skip-all", action="store_true",
                   help="skip research and LLM drafting; use deck content and keep final script approval")
    p.add_argument("--slide-gap", type=float, default=1.2, help="seconds of silence between slides")
    p.add_argument("--pause", type=float, default=0.8, help="seconds for each [pause] marker")
    p.add_argument("-o", "--output", help="output MP3 path")
    add_tts_opts(p)

    p = sub.add_parser("extract", help="stage 1: extract slide text + notes"); add_proj(p)
    p = sub.add_parser("research", help="stage 2: academic + real-world enrichment"); add_proj(p)
    p.add_argument("--audience", default="")
    p = sub.add_parser("script", help="stage 3: draft and approve the narration script"); add_proj(p)
    p.add_argument("--audience", default="")
    p.add_argument("--minutes", type=float)
    p.add_argument("--batch-size", type=int, default=15)
    p.add_argument("--skip-all", action="store_true",
                   help="skip LLM drafting and saved research; use deck content and keep final approval")
    p = sub.add_parser("tts", help="stage 4: synthesize MP3 from script.md"); add_proj(p)
    p.add_argument("--slide-gap", type=float, default=1.2)
    p.add_argument("--pause", type=float, default=0.8)
    p.add_argument("-o", "--output")
    add_tts_opts(p)
    p = sub.add_parser("voice-test", help="play a short sample of a voice")
    add_tts_opts(p)
    p.add_argument("-o", "--output", default="voice_sample.mp3")

    args = ap.parse_args()
    if args.cmd == "ui":
        raise SystemExit(subprocess.call([
            sys.executable, "-m", "streamlit", "run", str(Path(__file__).with_name("deck2audio_ui.py")),
            "--server.address", "127.0.0.1", "--server.headless", "true",
            "--browser.gatherUsageStats", "false",
        ], cwd=Path(__file__).resolve().parent))
    if args.cmd == "voice-test":
        voice = (args.macos_voice if args.backend == "macos" else
                 args.kokoro_voice if args.backend == "kokoro" else args.voice)
        asyncio.run(preview_voice(voice, args.rate, args.pitch, Path(args.output), backend=args.backend))
        return

    target = Path(args.target).expanduser()
    if target.suffix.lower() == ".pptx":
        if not target.exists():
            sys.exit(f"File not found: {target}")
        proj = project_dir_for(target)
    else:
        proj = target
    state = load_state(proj)

    try:
        if args.cmd == "extract" or (args.cmd == "run" and not state.get("extracted")):
            if target.suffix.lower() != ".pptx":
                sys.exit("Pass the .pptx file to extract.")
            extract(target, proj)
        skip_all = args.skip_all if args.cmd in ("script", "run") else False
        if args.cmd == "run" and not skip_all:
            skip_all = choose_skip_all()
        if args.cmd == "research" or (args.cmd == "run" and not skip_all):
            research(proj, args.audience)
        if args.cmd in ("script", "run"):
            script(proj, args.batch_size, args.audience, args.minutes,
                   skip_all=skip_all if args.cmd == "run" or skip_all else None)
        if args.cmd in ("tts", "run"):
            if args.cmd == "run" and load_state(proj).get("script") != "approved":
                print("Script not approved yet — stopping before audio. Re-run to continue.")
                return
            tts(proj, args)
    except (KeyboardInterrupt, EOFError):
        print("\nInterrupted — progress is saved; re-run the same command to resume.")


if __name__ == "__main__":
    main()
