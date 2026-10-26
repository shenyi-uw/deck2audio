import asyncio
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import aiohttp

from pptx import Presentation
from pptx.shapes.autoshape import Shape
from streamlit.testing.v1 import AppTest

import browser_workflow as workflow
import deck2audio as core


class ProjectTestCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.proj = Path(self.temp.name) / "talk_narration"
        self.proj.mkdir()
        self.slides = [
            {"number": 1, "title": "Early idea", "body": ["A key principle"], "notes": "First notes."},
            {"number": 2, "title": "Middle idea", "body": ["Everyday context"], "notes": ""},
            {"number": 3, "title": "Next steps", "body": ["Act thoughtfully"], "notes": "Last notes."},
        ]
        (self.proj / "slides.json").write_text(json.dumps(self.slides))
        core.save_state(self.proj, {"deck": "/example/talk.pptx", "extracted": True})

    def batch_text(self, numbers):
        return "\n\n".join(f"## Slide {n}\n\nNarration for slide {n}." for n in numbers)


class PromptGuidanceTests(ProjectTestCase):
    def test_final_batch_summarizes_entire_deck_with_context(self):
        prompt = core.script_prompt(self.slides, 2, 1, "high school students", 4.0,
                                    {2: "Previous narration"}, "")
        for instruction in (
            "END of the LAST slide", "ENTIRE presentation", "whether or not the deck contains a summary slide",
            "inside the last slide's section", "FULL DECK CONTEXT", "Early idea",
            "First notes.", "Previous narration", "high school students", "Wise, charismatic",
            "gentle, appropriate humor", "Fact-check claims", "speaker notes",
            "NOT to critique or review", "Do not repeat known false", "cannot verify",
            "relatable real-world examples", "no separate", "correlation from",
        ):
            self.assertIn(instruction, prompt)
        self.assertIn("Include EVERY slide from 3 to 3", prompt)

    def test_first_batch_does_not_conclude(self):
        prompt = core.script_prompt(self.slides, 2, 0, "graduate students", None, {}, "")
        self.assertIn("Do NOT conclude or summarize the whole talk yet", prompt)
        self.assertNotIn("FULL DECK CONTEXT", prompt)
        self.assertNotIn("Next steps", prompt)

    def test_later_batch_retains_all_earlier_examples_without_truncation(self):
        first = "A distinctive example from the first slide. " + "Additional narration. " * 100
        prompt = core.script_prompt(self.slides, 1, 2, "", None,
                                    {1: first, 2: "An entirely different example."}, "")
        self.assertIn("ACCEPTED EARLIER NARRATION", prompt)
        self.assertIn(first, prompt)
        self.assertIn("An entirely different example.", prompt)
        self.assertIn("do not rewrite or repeat", prompt)
        self.assertIn("Include EVERY slide from 3 to 3", prompt)
        self.assertIn("Use each case study, story, or analogy once", prompt)
        first_prompt = core.script_prompt(self.slides, 1, 0, "", None, {}, "")
        self.assertNotIn("ACCEPTED EARLIER NARRATION", first_prompt)

    def test_study_attribution_requires_finding_and_verified_context(self):
        prompt = core.script_prompt(self.slides, 15, 0, "", None, {}, "")
        for text in (
            "only when describing a specific study's finding",
            "Do not attach a researcher's name to an ordinary statement",
            "lead researcher", "institution at the time of the study", "year of",
            "publication", "omit that specific study reference",
            "commonly known", "public or historical figure", "current affiliation",
        ):
            self.assertIn(text, prompt)
        self.assertNotIn("As Daniel Kahneman showed", prompt)
        research = core.research_prompt("academic", self.slides)
        for text in ("lead researcher", "institution at the time", "year of publication",
                     "mark unknown details", "what the study actually found"):
            self.assertIn(text, research)

    def test_personable_podcast_style_without_invented_experience(self):
        prompt = core.script_prompt(self.slides, 15, 0, "", None, {}, "")
        for text in ('one-person podcast', 'one listener as "you"', 'use "I" naturally',
                     "Do not invent personal experiences", "brief reflective",
                     "gentle, appropriate humor"):
            self.assertIn(text, prompt)

    def test_research_examples_plan_avoids_repetition(self):
        prompt = core.research_prompt("examples", self.slides, "professionals")
        self.assertIn("Plan examples across the ENTIRE presentation", prompt)
        self.assertIn("best-fitting slide", prompt)
        self.assertIn("Do not repeat an example already supplied in the deck", prompt)

    def test_single_batch_requires_closing_summary(self):
        prompt = core.script_prompt(self.slides, 15, 0, "", None, {}, "")
        self.assertIn("ENTIRE presentation", prompt)
        self.assertIn("Early idea", prompt)
        self.assertIn("Next steps", prompt)
        self.slides[-1]["title"] = "Summary"
        prompt = core.script_prompt(self.slides, 15, 0, "", None, {}, "")
        self.assertIn("whether or not the deck contains a summary slide", prompt)

    def test_ui_launcher_is_local_headless_and_disables_telemetry(self):
        with patch("sys.argv", ["deck2audio.py", "ui"]), \
                patch("deck2audio.subprocess.call", return_value=0) as launch, \
                self.assertRaises(SystemExit) as exit_result:
            core.main()
        self.assertEqual(exit_result.exception.code, 0)
        command = launch.call_args.args[0]
        self.assertEqual(command[command.index("--server.address") + 1], "127.0.0.1")
        self.assertEqual(command[command.index("--server.headless") + 1], "true")
        self.assertEqual(command[command.index("--browser.gatherUsageStats") + 1], "false")
        self.assertEqual(launch.call_args.kwargs["cwd"], Path(core.__file__).resolve().parent)

    def test_research_audience_is_shared(self):
        for key in core.RESEARCH_PROMPTS:
            self.assertIn("K-12 students - grade 5", core.research_prompt(key, self.slides, "K-12 students - grade 5"))
        self.assertIn("relatable", core.research_prompt("examples", self.slides))

    def test_invalid_prompt_settings_raise(self):
        for size, index, minutes in ((0, 0, None), (1, 5, None), (2, 0, -1)):
            with self.subTest(size=size, index=index, minutes=minutes), self.assertRaises(ValueError):
                core.script_prompt(self.slides, size, index, "", minutes, {}, "")

    def test_validation_rejects_bad_shapes(self):
        samples = [
            ("## Slide 1\n\nOnly one.", "Missing"),
            (self.batch_text([1, 2, 3, 4]), "Unexpected"),
            ("## Slide 1\n\n[PAUSE]\n\n" + self.batch_text([2, 3]), "Empty"),
            (self.batch_text([1, 2, 2, 3]), "Duplicate"),
            (self.batch_text([2, 1, 3]), "deck order"),
            ("Here is your script\n" + self.batch_text([1, 2, 3]), "before"),
        ]
        for text, message in samples:
            with self.subTest(message=message):
                self.assertTrue(any(message in p for p in core.validate_script(text, {1, 2, 3})))
        self.assertEqual(core.validate_script(self.batch_text([1, 2, 3]), {1, 2, 3}), [])


class BrowserWorkflowTests(ProjectTestCase):
    def test_upload_extracts_and_resumes_same_content(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        assert title is not None
        title.text = "Uploaded deck"
        body = slide.placeholders[1]
        assert isinstance(body, Shape)
        body.text = "Body text"
        notes = slide.notes_slide.notes_text_frame
        assert notes is not None
        notes.text = "Speaker notes"
        buf = io.BytesIO()
        prs.save(buf)
        root = Path(self.temp.name) / "uploads"
        proj = workflow.upload_project(root, "../deck.pptx", buf.getvalue())
        self.assertEqual(core.load_slides(proj)[0]["title"], "Uploaded deck")
        self.assertEqual(core.load_slides(proj)[0]["notes"], "Speaker notes")
        core.save_state(proj, {"script": "approved"})
        self.assertEqual(workflow.upload_project(root, "../deck.pptx", buf.getvalue()), proj)
        self.assertEqual(core.load_state(proj)["script"], "approved")

    def test_upload_rejects_wrong_extension_and_empty(self):
        for name, data in (("talk.txt", b"content"), ("talk.pptx", b"")):
            with self.subTest(name=name), self.assertRaises(ValueError):
                workflow.upload_project(Path(self.temp.name), name, data)

    def test_skip_all_never_reads_research_or_calls_llm(self):
        with patch("deck2audio.approved_research", side_effect=AssertionError("research read")), \
                patch("deck2audio.human_in_the_loop", side_effect=AssertionError("LLM called")):
            workflow.start_pipeline(self.proj, workflow.PipelineConfig(skip_llm=True))
        self.assertEqual(core.load_state(self.proj)["ui"]["stage"], "review")
        self.assertEqual(workflow.pending_research(self.proj), [])
        text = (self.proj / "script.draft.md").read_text()
        self.assertIn("First notes.", text)
        self.assertIn("Everyday context", text)
        self.assertFalse(workflow.is_approved(self.proj, text))

    def test_research_approval_and_skip_advance(self):
        workflow.start_pipeline(self.proj, workflow.PipelineConfig(audience="undergraduate students"))
        self.assertEqual(workflow.pending_research(self.proj), ["academic", "examples"])
        with self.assertRaisesRegex(ValueError, "nonempty"):
            workflow.accept_research(self.proj, "academic", "")
        workflow.accept_research(self.proj, "academic", "Verified research")
        self.assertIn("Verified research", core.approved_research(self.proj))
        workflow.accept_research(self.proj, "examples", None)
        self.assertEqual(core.load_state(self.proj)["ui"]["stage"], "drafting")
        self.assertEqual(core.load_state(self.proj)["research"]["examples"], "skipped")
        self.assertIn("undergraduate students", (self.proj / "research/academic.prompt.txt").read_text())

    def test_batches_persist_and_final_prompt_sees_prior_deck(self):
        workflow.start_pipeline(self.proj, workflow.PipelineConfig(batch_size=2, academic=False, examples=False))
        with self.assertRaisesRegex(ValueError, "Missing"):
            workflow.accept_batch(self.proj, self.batch_text([1]))
        self.assertEqual(core.load_state(self.proj)["ui"]["next_batch"], 0)
        workflow.accept_batch(self.proj, self.batch_text([1, 2]))
        self.assertEqual(workflow.current_batch(self.proj)[0]["number"], 3)
        self.assertIn("Narration for slide 2", workflow.current_prompt(self.proj))
        self.assertIn("Narration for slide 1", workflow.current_prompt(self.proj))
        self.assertIn("ACCEPTED EARLIER NARRATION", workflow.current_prompt(self.proj))
        self.assertIn("Early idea", workflow.current_prompt(self.proj))
        workflow.accept_batch(self.proj, self.batch_text([3]))
        self.assertEqual(core.load_state(self.proj)["ui"]["stage"], "review")
        self.assertEqual(set(core.parse_script((self.proj / "script.draft.md").read_text())), {1, 2, 3})
        self.assertFalse((self.proj / "script.md").exists())

    def test_edit_or_restart_revokes_approval_and_audio(self):
        workflow.start_pipeline(self.proj, workflow.PipelineConfig(skip_llm=True))
        text = self.batch_text([1, 2, 3])
        workflow.save_review(self.proj, text, approved=True)
        self.assertTrue(workflow.is_approved(self.proj, text))
        state = core.load_state(self.proj)
        state.update({"audio": "old.mp3", "audio_source": "old"})
        core.save_state(self.proj, state)
        workflow.save_review(self.proj, "Incomplete edit", approved=False)
        self.assertFalse(workflow.is_approved(self.proj, text))
        self.assertNotIn("audio", core.load_state(self.proj))
        with self.assertRaises(ValueError):
            workflow.save_review(self.proj, "Incomplete edit", approved=True)
        workflow.save_review(self.proj, text, approved=True)
        workflow.start_pipeline(self.proj, workflow.PipelineConfig(academic=False, examples=False))
        self.assertEqual(core.load_state(self.proj)["script"], "draft")

    def test_audio_signature_tracks_each_setting(self):
        args = ["script", "voice", "+0%", "+0Hz", 1.2, 0.8]
        signature = workflow.audio_signature(*args)
        for index, replacement in enumerate(("edited", "other voice", "+10%", "+2Hz", 0.0, 0.0)):
            changed = args.copy()
            changed[index] = replacement
            self.assertNotEqual(signature, workflow.audio_signature(*changed))
        self.assertEqual(signature, workflow.audio_signature(*args, backend="edge"))
        self.assertNotEqual(signature, workflow.audio_signature(*args, backend="macos"))

    def test_shared_audio_generation_is_noninteractive_and_zero_pause_is_zero(self):
        audio = b"\xff\xf3\x80\xc4test"
        output = self.proj / "output.mp3"
        updates = []
        with patch("deck2audio.synth_chunk", new=AsyncMock(return_value=audio)) as synth, \
                patch("deck2audio.ask", side_effect=AssertionError("CLI input requested")), \
                patch("deck2audio.silent_mp3", side_effect=AssertionError("Zero pause inserted")):
            asyncio.run(core.generate_audio(self.proj, {1: "First [pause] next", 2: "Second"},
                                            "voice", "+0%", "+0Hz", 0, 0, output,
                                            progress=lambda done, total: updates.append((done, total))))
        self.assertEqual(synth.await_count, 3)
        self.assertEqual(output.read_bytes(), audio * 3)
        self.assertEqual((self.proj / "audio/slide_001.mp3").read_bytes(), audio * 2)
        self.assertEqual(core.load_state(self.proj)["audio"], str(output))
        self.assertEqual(updates, [(1, 2), (2, 2)])

    def test_audio_failure_does_not_record_success(self):
        output = self.proj / "output.mp3"
        with patch("deck2audio.synth_chunk", new=AsyncMock(side_effect=RuntimeError("Offline"))):
            with self.assertRaisesRegex(RuntimeError, "Offline"):
                asyncio.run(core.generate_audio(self.proj, {1: "First"}, "voice", "+0%", "+0Hz", 0, 0, output))
        self.assertFalse(output.exists())
        self.assertNotIn("audio", core.load_state(self.proj))


class BrowserInterfaceTests(ProjectTestCase):
    def app(self):
        app = AppTest.from_file(str(Path(__file__).with_name("deck2audio_ui.py")), default_timeout=15)
        app.session_state["active_project"] = str(self.proj)
        app.run()
        self.assertEqual(list(app.exception), [])
        return app

    def button(self, app, label):
        return next(button for button in app.button if button.label == label)

    def textarea(self, app, label):
        return next(area for area in app.text_area if area.label == label)

    def textarea_value(self, app, label) -> str:
        value = self.textarea(app, label).value
        assert value is not None
        return value

    def test_initial_page_has_pipeline_audience_and_voice_controls(self):
        app = self.app()
        self.assertEqual(app.title[0].value, "deck2audio")
        self.assertTrue(any("Applied Hearing Science Lab" in caption.value for caption in app.caption))
        self.assertIn("Intended audience", [box.label for box in app.selectbox])
        self.assertIn("Narrator voice", [box.label for box in app.selectbox])
        self.assertIn("Speaking rate (%)", [slider.label for slider in app.slider])
        self.assertIn("Pitch adjustment (Hz)", [slider.label for slider in app.slider])
        self.assertTrue(self.button(app, "Open uploaded deck").disabled)

    def test_lab_branding_assets_are_local_and_theme_matches_reference_palette(self):
        import tomllib

        root = Path(core.__file__).resolve().parent
        theme = tomllib.loads((root / ".streamlit/config.toml").read_text())["theme"]
        self.assertEqual(theme["base"], "dark")
        self.assertEqual(theme["primaryColor"], "#EFC563")
        self.assertEqual(theme["backgroundColor"], "#070D1B")
        self.assertEqual(theme["textColor"], "#F2F3C0")
        css = (root / "ui_branding.css").read_text()
        from PIL import Image

        with Image.open(root / "ahslab_logo.png") as logo:
            self.assertEqual(logo.format, "PNG")
            self.assertEqual(logo.size, (74, 81))
            logo.verify()
        self.assertIn(".ahslab-logo", css)
        self.assertIn("#394c7d", css)
        self.assertIn("@media", css)
        self.assertIn(":focus-visible", css)
        self.assertNotIn("https://", css)
        self.assertNotIn("@import", css)

    def test_skip_all_form_review_approval_audio_and_edit(self):
        app = self.app()
        next(box for box in app.checkbox if box.label.startswith("Skip all")).check()
        self.button(app, "Start / restart pipeline").click().run()
        self.assertEqual(list(app.exception), [])
        self.assertIn("First notes.", self.textarea_value(app, "Edit the full narration script"))
        self.assertTrue(self.button(app, "Generate MP3 (online)").disabled)
        self.button(app, "Approve script for audio").click().run()
        self.assertFalse(self.button(app, "Generate MP3 (online)").disabled)
        with patch("deck2audio.synth_chunk", new=AsyncMock(return_value=b"\xff\xf3\x80\xc4test")):
            self.button(app, "Generate MP3 (online)").click().run()
        self.assertEqual(list(app.exception), [])
        self.assertTrue((self.proj / "narrated.mp3").exists())
        self.assertIn("audio_source", core.load_state(self.proj))
        self.textarea(app, "Edit the full narration script").set_value("Incomplete edited script").run()
        self.assertEqual(list(app.exception), [])
        self.assertTrue(self.button(app, "Generate MP3 (online)").disabled)
        self.assertTrue(self.button(app, "Approve script for audio").disabled)
        self.assertEqual(core.load_state(self.proj)["script"], "draft")
        self.assertNotIn("audio_source", core.load_state(self.proj))

    def test_llm_textboxes_batch_validation_and_resume(self):
        workflow.start_pipeline(self.proj, workflow.PipelineConfig(batch_size=2, academic=False, examples=False))
        app = self.app()
        self.assertIn("Fact-check", self.textarea_value(app, "Prompt to copy"))
        self.textarea(app, "Paste and edit the LLM script response").set_value(self.batch_text([1])).run()
        self.button(app, "Accept batch and continue").click().run()
        self.assertTrue(list(app.error))
        self.assertEqual(core.load_state(self.proj)["ui"]["next_batch"], 0)
        self.textarea(app, "Paste and edit the LLM script response").set_value(self.batch_text([1, 2])).run()
        self.button(app, "Accept batch and continue").click().run()
        self.assertEqual(list(app.exception), [])
        self.assertIn("ENTIRE presentation", self.textarea_value(app, "Prompt to copy"))
        resumed = self.app()
        self.assertIn("Early idea", self.textarea_value(resumed, "Prompt to copy"))
        self.textarea(resumed, "Paste and edit the LLM script response").set_value(self.batch_text([3])).run()
        self.button(resumed, "Accept batch and continue").click().run()
        self.assertEqual(list(resumed.exception), [])
        self.assertIn("Narration for slide 1", self.textarea_value(resumed, "Edit the full narration script"))

    def test_research_response_and_skip(self):
        workflow.start_pipeline(self.proj, workflow.PipelineConfig())
        app = self.app()
        self.textarea(app, "Paste and edit the LLM research response").set_value("Approved academic context").run()
        self.button(app, "Approve research response").click().run()
        self.assertEqual(list(app.exception), [])
        self.button(app, "Skip this research step").click().run()
        self.assertIn("Approved academic context", self.textarea_value(app, "Prompt to copy"))

    def test_voice_preview_failure_is_visible(self):
        app = self.app()
        with patch("deck2audio.synth_chunk", new=AsyncMock(side_effect=RuntimeError("Network unavailable"))):
            self.button(app, "Preview voice (online)").click().run()
        self.assertEqual(list(app.exception), [])
        self.assertTrue(any("Network unavailable" in error.value for error in app.error))

    def test_online_403_suggests_local_fallback_without_switching(self):
        with patch("macos_tts.available", return_value=True):
            app = self.app()
            error = aiohttp.WSServerHandshakeError(
                MagicMock(real_url="https://speech.platform.bing.com"), (), status=403,
                message="Invalid response status")
            with patch("deck2audio.synth_chunk", new=AsyncMock(side_effect=error)):
                self.button(app, "Preview voice (online)").click().run()
        self.assertEqual(list(app.exception), [])
        self.assertTrue(any("403" in item.value for item in app.error))
        self.assertTrue(any("emergency fallback" in item.value for item in app.info))
        self.assertEqual(core.load_state(self.proj)["ui_voice"]["backend"], "edge")

    def test_failed_regeneration_hides_previous_audio(self):
        workflow.start_pipeline(self.proj, workflow.PipelineConfig(skip_llm=True))
        text = (self.proj / "script.draft.md").read_text()
        workflow.save_review(self.proj, text, approved=True)
        state = core.load_state(self.proj)
        state.update({"audio": str(self.proj / "old.mp3"), "audio_source": "old"})
        core.save_state(self.proj, state)
        app = self.app()
        with patch("deck2audio.synth_chunk", new=AsyncMock(side_effect=RuntimeError("Offline"))):
            self.button(app, "Generate MP3 (online)").click().run()
        self.assertEqual(list(app.exception), [])
        self.assertTrue(any("Offline" in error.value for error in app.error))
        self.assertNotIn("audio_source", core.load_state(self.proj))
        self.assertNotIn("audio", core.load_state(self.proj))

    def test_voice_settings_persist_across_sessions(self):
        app = self.app()
        next(slider for slider in app.slider if slider.label == "Speaking rate (%)").set_value(10).run()
        resumed = self.app()
        self.assertEqual(next(slider for slider in resumed.slider if slider.label == "Speaking rate (%)").value, 10)

    def test_macos_fallback_is_explicit_persisted_and_requires_approval(self):
        workflow.start_pipeline(self.proj, workflow.PipelineConfig(skip_llm=True))
        with patch("macos_tts.available", return_value=True), \
                patch("macos_tts.list_voices", return_value={"Samantha": "en_US"}):
            app = self.app()
            backend = next(box for box in app.selectbox if box.label == "Speech backend")
            self.assertEqual(backend.value, "Edge voices (online)")
            backend.select("macOS voices (offline fallback)").run()
            self.assertEqual(list(app.exception), [])
            self.assertTrue(self.button(app, "Generate MP3 (offline)").disabled)
            self.assertNotIn("Pitch adjustment (Hz)", [slider.label for slider in app.slider])
            next(box for box in app.selectbox if box.label == "Installed macOS voice").select("Samantha").run()
            with patch("deck2audio.synth_chunk", new=AsyncMock(return_value=b"mp3")) as preview:
                self.button(app, "Preview voice (offline)").click().run()
            preview.assert_awaited_once_with(core.VOICE_PREVIEW_TEXT, "Samantha", "+0%", "+0Hz", backend="macos")
            self.button(app, "Approve script for audio").click().run()
            with patch("macos_tts.validate_settings"), \
                    patch("macos_tts.synth_pcm", new=AsyncMock(return_value=b"\x01\x00" * 12)), \
                    patch("pcm_audio.encode_mp3", return_value=b"mp3"), \
                    patch("deck2audio.synth_chunk", side_effect=AssertionError("Edge called")):
                self.button(app, "Generate MP3 (offline)").click().run()
            self.assertEqual(list(app.exception), [])
            state = core.load_state(self.proj)
            self.assertEqual(state["ui_voice"]["backend"], "macos")
            self.assertEqual(state["ui_voice"]["macos_voice"], "Samantha")
            self.assertEqual(state["audio_source"], workflow.audio_signature(
                (self.proj / "script.md").read_text(), "Samantha", "+0%", "+0Hz", 1.2, 0.8, backend="macos"))
            resumed = self.app()
            self.assertEqual(next(box for box in resumed.selectbox if box.label == "Speech backend").value,
                             "macOS voices (offline fallback)")
            next(box for box in app.selectbox if box.label == "Speech backend").select("Edge voices (online)").run()
            self.assertTrue(any("settings changed" in info.value for info in app.info))

    def test_non_macos_offers_all_cross_platform_backends(self):
        with patch("macos_tts.available", return_value=False):
            app = self.app()
        backend = next(box for box in app.selectbox if box.label == "Speech backend")
        self.assertEqual(list(backend.options), [
            "Edge voices (online)", "Kokoro voices (offline)",
            "VibeVoice voices (offline, experimental)", "My voice - Chatterbox (offline, experimental)"])

    def test_kokoro_preview_generation_persistence_and_backend_switch(self):
        workflow.start_pipeline(self.proj, workflow.PipelineConfig(skip_llm=True))
        app = self.app()
        next(box for box in app.selectbox if box.label == "Speech backend").select("Kokoro voices (offline)").run()
        self.assertTrue(self.button(app, "Generate MP3 (offline)").disabled)
        self.assertNotIn("Pitch adjustment (Hz)", [slider.label for slider in app.slider])
        next(box for box in app.selectbox if box.label == "Kokoro voice").select("af_bella").run()
        with patch("deck2audio.synth_chunk", new=AsyncMock(return_value=b"mp3")) as preview:
            self.button(app, "Preview voice (offline)").click().run()
        preview.assert_awaited_once_with(core.VOICE_PREVIEW_TEXT, "af_bella", "+0%", "+0Hz", backend="kokoro")
        self.button(app, "Approve script for audio").click().run()
        with patch("kokoro_tts.synth_pcm", new=AsyncMock(return_value=b"\x01\x00" * 12)), \
                patch("pcm_audio.encode_mp3", return_value=b"mp3"), \
                patch("edge_tts.Communicate", side_effect=AssertionError("Edge called")):
            self.button(app, "Generate MP3 (offline)").click().run()
        self.assertEqual(list(app.exception), [])
        state = core.load_state(self.proj)
        self.assertEqual(state["ui_voice"]["kokoro_voice"], "af_bella")
        self.assertEqual(state["ui_voice"]["backend"], "kokoro")
        self.assertEqual(state["audio_source"], workflow.audio_signature(
            (self.proj / "script.md").read_text(), "af_bella", "+0%", "+0Hz", 1.2, 0.8, backend="kokoro"))
        resumed = self.app()
        self.assertEqual(next(box for box in resumed.selectbox if box.label == "Speech backend").value,
                         "Kokoro voices (offline)")
        self.assertEqual(next(box for box in resumed.selectbox if box.label == "Kokoro voice").value, "af_bella")
        next(box for box in app.selectbox if box.label == "Speech backend").select("Edge voices (online)").run()
        self.assertTrue(any("settings changed" in info.value for info in app.info))

    def test_kokoro_setup_failure_is_visible(self):
        app = self.app()
        next(box for box in app.selectbox if box.label == "Speech backend").select("Kokoro voices (offline)").run()
        with patch("kokoro_tts.synth_pcm", new=AsyncMock(side_effect=RuntimeError("Kokoro model files are missing"))):
            self.button(app, "Preview voice (offline)").click().run()
        self.assertEqual(list(app.exception), [])
        self.assertTrue(any("model files are missing" in error.value for error in app.error))

    def test_vibevoice_preview_generation_persistence_and_natural_speed(self):
        import neural_tts

        workflow.start_pipeline(self.proj, workflow.PipelineConfig(skip_llm=True))
        app = self.app()
        next(slider for slider in app.slider if slider.label == "Speaking rate (%)").set_value(10).run()
        next(box for box in app.selectbox if box.label == "Speech backend").select(
            "VibeVoice voices (offline, experimental)").run()
        self.assertEqual(list(app.exception), [])
        self.assertNotIn("Speaking rate (%)", [slider.label for slider in app.slider])
        self.assertNotIn("Pitch adjustment (Hz)", [slider.label for slider in app.slider])
        self.assertTrue(self.button(app, "Generate MP3 (offline)").disabled)
        next(box for box in app.selectbox if box.label == "VibeVoice narrator").select("Grace").run()
        with patch("deck2audio.synth_chunk", new=AsyncMock(return_value=b"mp3")) as preview:
            self.button(app, "Preview voice (offline)").click().run()
        preview.assert_awaited_once_with(core.VOICE_PREVIEW_TEXT, "Grace", "+0%", "+0Hz", backend="vibevoice")
        self.button(app, "Approve script for audio").click().run()
        with patch.object(neural_tts.Worker, "synthesize", return_value=b"\x01\0" * 12), \
                patch("pcm_audio.encode_mp3", return_value=b"mp3"):
            self.button(app, "Generate MP3 (offline)").click().run()
        self.assertEqual(list(app.exception), [])
        state = core.load_state(self.proj)
        self.assertEqual(state["ui_voice"]["vibevoice_voice"], "Grace")
        self.assertEqual(state["ui_voice"]["rate"], 10)
        self.assertEqual(state["audio_source"], workflow.audio_signature(
            (self.proj / "script.md").read_text(), "Grace", "+0%", "+0Hz", 1.2, 0.8, backend="vibevoice"))
        resumed = self.app()
        self.assertEqual(next(box for box in resumed.selectbox if box.label == "VibeVoice narrator").value, "Grace")
        next(box for box in app.selectbox if box.label == "Speech backend").select("Edge voices (online)").run()
        self.assertTrue(any("settings changed" in info.value for info in app.info))

    def test_microphone_recording_requires_consent_saves_previews_and_deletes(self):
        from test_neural_tts import recording

        workflow.start_pipeline(self.proj, workflow.PipelineConfig(skip_llm=True))
        app = self.app()
        next(box for box in app.selectbox if box.label == "Speech backend").select(
            "My voice - Chatterbox (offline, experimental)").run()
        self.assertEqual(list(app.exception), [])
        self.assertTrue(self.button(app, "Save voice reference").disabled)
        self.assertTrue(self.button(app, "Preview voice (offline)").disabled)
        self.assertEqual(len(app.get("audio_input")), 1)
        source = MagicMock()
        source.getvalue.return_value = recording()
        with patch("streamlit.audio_input", return_value=source):
            app.run()
            self.assertTrue(self.button(app, "Save voice reference").disabled)
            next(box for box in app.checkbox if box.label.startswith("This is my voice")).check().run()
            self.button(app, "Save voice reference").click().run()
        self.assertEqual(list(app.exception), [])
        references = list((self.proj / "voices").glob("*.wav"))
        self.assertEqual(len(references), 1)
        self.assertEqual(references[0].read_bytes(), recording())
        voice = str(references[0].resolve())
        self.assertEqual(core.load_state(self.proj)["ui_voice"]["voice_reference"], voice)
        with patch("deck2audio.synth_chunk", new=AsyncMock(return_value=b"mp3")) as preview:
            self.button(app, "Preview voice (offline)").click().run()
        preview.assert_awaited_once_with(core.VOICE_PREVIEW_TEXT, voice, "+0%", "+0Hz", backend="chatterbox")
        self.button(app, "Approve script for audio").click().run()
        import neural_tts
        with patch.object(neural_tts.Worker, "synthesize", return_value=b"\x01\0" * 12), \
                patch("pcm_audio.encode_mp3", return_value=b"mp3"):
            self.button(app, "Generate MP3 (offline)").click().run()
        self.assertEqual(list(app.exception), [])
        resumed = self.app()
        self.assertEqual(next(box for box in resumed.selectbox if box.label == "Saved voice reference").value, voice)
        self.button(app, "Delete selected voice reference").click().run()
        self.assertEqual(list(app.exception), [])
        self.assertFalse(references[0].exists())
        self.assertEqual(core.load_state(self.proj)["ui_voice"]["voice_reference"], "")
        self.assertTrue(self.button(app, "Generate MP3 (offline)").disabled)
        self.assertNotIn("audio_source", core.load_state(self.proj))

    def test_invalid_microphone_recording_reports_error_without_saving(self):
        app = self.app()
        next(box for box in app.selectbox if box.label == "Speech backend").select(
            "My voice - Chatterbox (offline, experimental)").run()
        source = MagicMock()
        source.getvalue.return_value = b"invalid wav"
        with patch("streamlit.audio_input", return_value=source):
            next(box for box in app.checkbox if box.label.startswith("This is my voice")).check().run()
            self.button(app, "Save voice reference").click().run()
        self.assertEqual(list(app.exception), [])
        self.assertTrue(any("valid WAV" in error.value for error in app.error))
        self.assertFalse((self.proj / "voices").exists())


if __name__ == "__main__":
    unittest.main()
